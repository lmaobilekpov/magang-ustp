from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from dokumen.models import Karyawan


class Command(BaseCommand):
    help = "Import data karyawan dari Excel."

    def add_arguments(self, parser):
        parser.add_argument(
            "file_path",
            type=str,
            help="Path ke file Excel yang berisi data karyawan.",
        )
        parser.add_argument(
            "--sheet",
            default="karyawa",
            help="Nama sheet Excel yang berisi data karyawan.",
        )

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError:
            raise CommandError(
                "Paket openpyxl belum terpasang. Jalankan: pip install openpyxl"
            )

        file_path = Path(options["file_path"])
        sheet_name = options["sheet"]

        if not file_path.exists():
            raise CommandError(f"File tidak ditemukan: {file_path}")

        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        except Exception as error:
            raise CommandError(f"File Excel tidak dapat dibaca: {error}")

        if sheet_name not in workbook.sheetnames:
            workbook.close()
            raise CommandError(
                f"Sheet '{sheet_name}' tidak ditemukan. Sheet tersedia: {', '.join(workbook.sheetnames)}"
            )

        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)

        try:
            headers = next(rows)
        except StopIteration:
            workbook.close()
            raise CommandError("Sheet karyawan kosong.")

        header_map = {
            str(header).strip().upper(): index
            for index, header in enumerate(headers)
            if header is not None
        }

        required_headers = ["EMPCODE", "EMPNAME", "JABATAN", "DOB"]
        missing_headers = [header for header in required_headers if header not in header_map]

        if missing_headers:
            workbook.close()
            raise CommandError(
                "Kolom wajib tidak ditemukan: " + ", ".join(missing_headers)
            )

        imported_count = 0
        updated_count = 0
        skipped_count = 0

        with transaction.atomic():
            for row_number, row in enumerate(rows, start=2):
                raw_code = row[header_map["EMPCODE"]]
                raw_name = row[header_map["EMPNAME"]]
                raw_position = row[header_map["JABATAN"]]
                raw_dob = row[header_map["DOB"]]

                code = str(raw_code).strip() if raw_code is not None else ""
                name = str(raw_name).strip() if raw_name is not None else ""
                position = str(raw_position).strip() if raw_position is not None else ""

                if not code or not name or not position or raw_dob is None:
                    skipped_count += 1
                    self.stdout.write(
                        self.style.WARNING(
                            f"Baris {row_number} dilewati karena kode, nama, jabatan, atau DOB kosong."
                        )
                    )
                    continue

                if hasattr(raw_dob, "date"):
                    dob = raw_dob.date()
                else:
                    try:
                        from datetime import datetime
                        dob = datetime.strptime(str(raw_dob).strip(), "%d-%m-%Y").date()
                    except (TypeError, ValueError):
                        skipped_count += 1
                        self.stdout.write(
                            self.style.WARNING(
                                f"Baris {row_number} dilewati karena format DOB tidak valid: {raw_dob}"
                            )
                        )
                        continue

                karyawan, created = Karyawan.objects.update_or_create(
                    kode_karyawan=code,
                    defaults={
                        "nama_lengkap": name,
                        "jabatan": position,
                        "tanggal_lahir": dob,
                    },
                )

                if created:
                    imported_count += 1
                else:
                    updated_count += 1

        workbook.close()

        self.stdout.write(self.style.SUCCESS("Import karyawan selesai."))
        self.stdout.write(f"Data baru: {imported_count}")
        self.stdout.write(f"Data diperbarui: {updated_count}")
        self.stdout.write(f"Baris dilewati: {skipped_count}")
        self.stdout.write(f"Total karyawan di database: {Karyawan.objects.count()}")
